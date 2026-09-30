"""Importação da planilha antiga usando um arquivo fictício com a mesma estrutura."""

from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from rfq_control.modelos import StatusRFQ
from rfq_control.servicos.exportacao import exportar_rfqs
from rfq_control.servicos.importacao_legado import importar_planilha
from rfq_control.servicos.leitor_xlsx import ErroPlanilha, LeitorXlsx, data_excel

CABECALHO_CONTROLE = [
    "RFQ Nº", "DATA", "SOLICITANTE", "CLIENTE", "PROJETO", "FORNECEDOR ", "OP - Ref", "CDC - Ref",
    "PART DESCRIPTION", "UM / QTY", "Volume \nannual", "Please confirm 120% capacity", "JUST enter", "", "",
    "RFQ Recebida",
]


def _serial(dia: date) -> int:
    return (dia - date(1899, 12, 30)).days


@pytest.fixture
def planilha_antiga(tmp_path):
    livro = Workbook()
    controle = livro.active
    controle.title = "Controle"
    controle.append(CABECALHO_CONTROLE)
    d1, d2 = _serial(date(2026, 3, 19)), _serial(date(2026, 3, 20))
    itens = [("PECA A", "CDC PECA A", "PC"), ("PECA B", "CDC PECA B", 154000)]
    # mesmo pacote enviado a dois fornecedores (RFQ001 e RFQ002)
    for numero, fornecedor, dia in (("RFQ2026001", "FORN UM ", d1), ("RFQ2026002", "Forn Dois", d2)):
        for ref_op, ref_cdc, um in itens:
            controle.append([numero, dia, "SOLIC TESTE ", "CLI A", "PROJ ALFA", fornecedor, ref_op, ref_cdc, None, um, "NA"])
    # RFQ com projeto fora do cadastro e dois fornecedores (vira RFQ2026003 e RFQ2026003-2)
    controle.append(["RFQ2026003", d2, "SOLIC TESTE ", "", "PROJ NOVO ", "FORN UM", "X", "CDC X", "desc", "EA"])
    controle.append(["RFQ2026003", d2, "SOLIC TESTE ", "", "PROJ NOVO ", "Forn Tres", "Y", "CDC Y", None, "EA"])
    # linhas "fantasma" com só a coluna D preenchida (como na planilha real)
    for _ in range(50):
        controle.append([None, None, None, ""])

    projetos = livro.create_sheet("CadastroProjetos")
    projetos.append(["PLANTA OP ", "Cliente", "Projeto", "SOP ", "Lifetime (Years)", None, None, None, None, "Requester"])
    projetos.append(["Planta Sul", "CLI A", "PROJ ALFA ", None, 6, None, "Planta", None, None, "SOLIC TESTE "])
    projetos.append(["Planta Sul", "CLI B", "PROJ BETA", None, None])

    contatos = livro.create_sheet("CadastroContatos")
    contatos.append(["FORNECEDOR ", "RESPONSÁVEL", " EMAIL (To)", "EMAIL (Cc)", "BUYER EMAIL"])
    contatos.append(["FORN UM", "Carlos", "carlos@um.com; vendas@um.com", "gerente@um.com", "compras@exemplo.com"])
    contatos.append(["Forn Dois", "fornecedor", "contato@dois.com", "email-invalido", None])
    contatos.append(["forn um ", "Repetido", "outro@um.com"])

    feriados = livro.create_sheet("Feriados")
    feriados.append([_serial(date(2025, 12, 25))])
    feriados.append([_serial(date(2026, 3, 23))])  # feriado local fictício

    caminho = tmp_path / "RFQ_Controle_teste.xlsx"
    livro.save(caminho)
    return caminho


def test_leitor_para_apos_linhas_vazias(planilha_antiga):
    with LeitorXlsx(planilha_antiga) as leitor:
        assert {"Controle", "CadastroProjetos", "CadastroContatos", "Feriados"} <= set(leitor.abas)
        linhas = list(leitor.linhas("Controle", inicio=2, parar_apos_vazias=10))
    assert len(linhas) == 6
    numero, valores = linhas[0]
    assert numero == 2 and valores["A"] == "RFQ2026001" and data_excel(valores["B"]) == date(2026, 3, 19)


def test_arquivo_invalido(tmp_path):
    arquivo = tmp_path / "nao_e_planilha.xlsm"
    arquivo.write_text("texto")
    with pytest.raises(ErroPlanilha):
        LeitorXlsx(arquivo)


def test_importacao_completa(base, planilha_antiga):
    relatorio = importar_planilha(base, planilha_antiga)

    assert relatorio.linhas_lidas == 6
    assert relatorio.rfqs_importadas == 4  # 001, 002, 003, 003-2
    assert relatorio.pacotes_criados == 3  # 001+002 compartilham o pacote
    assert relatorio.pasta_backup and relatorio.pasta_backup.exists()

    numeros = sorted(r.numero for r in base.rfqs)
    assert numeros == ["RFQ2026001", "RFQ2026002", "RFQ2026003", "RFQ2026003-2"]
    rfq1 = next(r for r in base.rfqs if r.numero == "RFQ2026001")
    rfq2 = next(r for r in base.rfqs if r.numero == "RFQ2026002")
    assert rfq1.pacote_id == rfq2.pacote_id
    assert rfq1.status == StatusRFQ.ENVIADA
    assert rfq1.data_envio == date(2026, 3, 19)
    # 19/03 (qui) + 4 dias úteis pulando o feriado de 23/03 → 26/03
    assert rfq1.prazo == date(2026, 3, 26)

    pacote = base.pacotes.obter(rfq1.pacote_id)
    assert pacote.data == date(2026, 3, 19)
    assert [(i.ref_op, i.unidade, i.quantidade, i.volume_anual) for i in pacote.itens] == [
        ("PECA A", "PC", None, None), ("PECA B", "", 154000, None)
    ]

    projetos = {p.nome: p for p in base.projetos}
    assert set(projetos) == {"PROJ ALFA", "PROJ BETA", "PROJ NOVO"}
    assert projetos["PROJ ALFA"].lifetime_anos == 6 and projetos["PROJ ALFA"].planta == "Planta Sul"

    fornecedores = {f.nome: f for f in base.fornecedores}
    assert set(fornecedores) == {"FORN UM", "Forn Dois", "Forn Tres"}
    um = fornecedores["FORN UM"]
    assert um.emails_para() == ["carlos@um.com", "vendas@um.com"]
    assert um.emails_copia() == ["gerente@um.com", "compras@exemplo.com"]
    assert um.nome_contato() == "Carlos"
    assert fornecedores["Forn Dois"].nome_contato() == ""  # "fornecedor" não é nome de pessoa

    assert [s.nome for s in base.solicitantes] == ["SOLIC TESTE"]
    assert date(2026, 3, 23) in base.datas_feriados()

    avisos = "\n".join(relatorio.avisos)
    assert "E-mail inválido ignorado para Forn Dois" in avisos
    assert "Fornecedor repetido" in avisos
    assert "PROJ NOVO" in avisos and "RFQ2026003-2" in avisos


def test_importacao_repetida_nao_duplica(base, planilha_antiga):
    importar_planilha(base, planilha_antiga)
    segunda = importar_planilha(base, planilha_antiga)
    assert segunda.rfqs_importadas == 0 and segunda.rfqs_existentes == 4
    assert segunda.pacotes_criados == 0 and segunda.fornecedores_criados == 0
    assert len(base.rfqs) == 4 and len(base.fornecedores) == 3


def test_importacao_sem_aba_controle(base, tmp_path):
    livro = Workbook()
    livro.active.title = "Outra"
    caminho = tmp_path / "outra.xlsx"
    livro.save(caminho)
    with pytest.raises(ErroPlanilha, match="aba 'Controle'"):
        importar_planilha(base, caminho)


def test_exportar_excel(base, planilha_antiga, tmp_path):
    importar_planilha(base, planilha_antiga)
    destino = exportar_rfqs(base, tmp_path / "saida" / "rfqs.xlsx", hoje=date(2026, 3, 20))
    livro = load_workbook(destino)
    assert livro.sheetnames == ["RFQs", "Itens"]
    rfqs = list(livro["RFQs"].iter_rows(min_row=2, values_only=True))
    assert len(rfqs) == 4 and rfqs[0][0] == "RFQ2026001" and rfqs[0][1] == "Enviada"
    assert livro["Itens"].max_row == 1 + 2 + 2 + 1 + 1
