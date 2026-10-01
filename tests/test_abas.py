"""Abas da base de dados e cadastro em massa (Excel/CSV), com dados fictícios."""

from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from rfq_control.modelos import Contato, Fornecedor, Idioma, StatusRFQ
from rfq_control.servicos import abas
from rfq_control.servicos.importacao_legado import importar_planilha
from rfq_control.servicos.rfq import ErroNegocio


def _csv(tmp_path, texto: str, codificacao: str = "utf-8-sig"):
    caminho = tmp_path / "dados.csv"
    caminho.write_bytes(texto.encode(codificacao))
    return caminho


def test_importar_fornecedores_csv_com_idioma(base, tmp_path):
    caminho = _csv(
        tmp_path,
        "Fornecedor;Responsável;E-mail (Para);E-mail (Cc);Idioma;Ativo\n"
        "Alfa Ltda;Carla;carla@alfa.com;compras@alfa.com;Espanhol;Sim\n"
        "Beta SA;;vendas@beta.com, outro@beta.com;;Inglês;\n"
        "Gama;Rui;rui(arroba)gama;;Português;Sim\n",
        "cp1252",
    )
    resultado = abas.importar_arquivo(base, abas.ABA_FORNECEDORES, caminho)

    assert resultado.criados == 2 and len(resultado.erros) == 1
    assert "Linha 4" in resultado.erros[0] and "E-mail inválido" in resultado.erros[0]
    assert resultado.pasta_backup and resultado.pasta_backup.exists()
    fornecedores = {f.nome: f for f in base.fornecedores}
    assert fornecedores["Alfa Ltda"].idioma == Idioma.ES
    assert fornecedores["Alfa Ltda"].nome_contato() == "Carla"
    assert fornecedores["Alfa Ltda"].emails_copia() == ["compras@alfa.com"]
    assert fornecedores["Beta SA"].idioma == Idioma.EN
    assert fornecedores["Beta SA"].emails_para() == ["vendas@beta.com", "outro@beta.com"]


def test_importacao_atualiza_pelo_nome_sem_perder_contatos(base, tmp_path):
    base.fornecedores.salvar(Fornecedor(
        nome="Alfa Ltda",
        contatos=[
            Contato(nome="Carla", email="carla@alfa.com", telefone="11 9999-0000"),
            Contato(nome="Diretor", email="diretor@alfa.com", copia=True),
        ],
    ))
    caminho = _csv(tmp_path, "FORNECEDOR,Idioma,Categoria\n  alfa   ltda ,Español,Ferramentaria\n")
    resultado = abas.importar_arquivo(base, abas.ABA_FORNECEDORES, caminho)

    assert (resultado.criados, resultado.atualizados) == (0, 1)
    [fornecedor] = base.fornecedores.todos()
    assert fornecedor.idioma == Idioma.ES and fornecedor.categoria == "Ferramentaria"
    # colunas de contato ausentes no arquivo: contatos mantidos
    assert [(c.nome, c.email, c.telefone) for c in fornecedor.contatos] == [
        ("Carla", "carla@alfa.com", "11 9999-0000"), ("Diretor", "diretor@alfa.com", ""),
    ]


def test_cabecalhos_da_planilha_antiga(base, tmp_path):
    livro = Workbook()
    aba = livro.active
    aba.title = "CadastroContatos"
    aba.append(["FORNECEDOR ", "RESPONSÁVEL", " EMAIL (To)", "EMAIL (Cc)", "BUYER EMAIL", "Coluna1"])
    aba.append(["Delta", "fornecedor", "a@delta.com", "b@delta.com", "comprador@empresa.com", None])
    caminho = tmp_path / "antiga.xlsx"
    livro.save(caminho)

    resultado = abas.importar_arquivo(base, abas.ABA_FORNECEDORES, caminho)
    assert resultado.criados == 1
    assert resultado.colunas_ignoradas == ["Coluna1"]
    [delta] = base.fornecedores.todos()
    assert delta.nome_contato() == ""  # "fornecedor" não é nome de pessoa
    assert delta.emails_copia() == ["b@delta.com", "comprador@empresa.com"]


def test_modelo_exportacao_e_reimportacao(base, tmp_path):
    modelo = abas.gerar_modelo(abas.ABA_FORNECEDORES, tmp_path / "modelo.xlsx")
    livro = load_workbook(modelo)
    planilha = livro[abas.ABA_FORNECEDORES.titulo]
    assert [c.value for c in planilha[1]] == [c.titulo for c in abas.ABA_FORNECEDORES.colunas]
    validacoes = [v.formula1 for v in planilha.data_validations.dataValidation]
    assert '"Português,Espanhol,Inglês"' in validacoes
    assert "Instruções" in livro.sheetnames

    planilha.append(["Épsilon", "Ana", "ana@epsilon.com", "", "Inglês", "Fixadores", "", "Sim", ""])
    livro.save(modelo)
    assert abas.importar_arquivo(base, abas.ABA_FORNECEDORES, modelo).criados == 1

    exportado = abas.exportar_aba(base, abas.ABA_FORNECEDORES, tmp_path / "exportado.xlsx")
    novamente = abas.importar_arquivo(base, abas.ABA_FORNECEDORES, exportado)
    assert (novamente.criados, novamente.atualizados, novamente.sem_alteracao) == (0, 0, 1)


def test_projetos_e_erros_de_tipo(base, tmp_path):
    caminho = _csv(
        tmp_path,
        "PLANTA OP;Cliente;Projeto;SOP;Lifetime (Years)\n"
        "Planta Sul;Cliente A;Projeto X;15/03/2027;6\n"
        "Planta Sul;Cliente B;Projeto Y;31/02/2027;\n"
        "Planta Sul;Cliente C;Projeto Z;;seis\n",
    )
    resultado = abas.importar_arquivo(base, abas.ABA_PROJETOS, caminho)
    assert resultado.criados == 1 and len(resultado.erros) == 2
    assert "SOP" in resultado.erros[0] and "Lifetime" in resultado.erros[1]
    [projeto] = base.projetos.todos()
    assert projeto.sop == date(2027, 3, 15) and projeto.lifetime_anos == 6


def test_feriados_sem_cabecalho_como_na_planilha_antiga(base, tmp_path):
    livro = Workbook()
    aba = livro.active
    aba.title = "Feriados"
    for dia in (date(2031, 1, 1), date(2031, 4, 21), date(2031, 1, 1)):
        aba.append([(dia - date(1899, 12, 30)).days])
    caminho = tmp_path / "feriados.xlsx"
    livro.save(caminho)
    antes = len(base.feriados)
    resultado = abas.importar_arquivo(base, abas.ABA_FERIADOS, caminho)
    assert resultado.criados == 2 and resultado.sem_alteracao == 1
    assert len(base.feriados) == antes + 2


def test_arquivo_sem_a_coluna_obrigatoria(base, tmp_path):
    caminho = _csv(tmp_path, "Coluna;Outra\n1;2\n")
    with pytest.raises(ErroNegocio, match="Não encontrei a coluna 'Fornecedor'"):
        abas.importar_arquivo(base, abas.ABA_FORNECEDORES, caminho)


def test_linha_da_grade(base):
    definicao = abas.ABA_FORNECEDORES
    with pytest.raises(ErroNegocio, match="Preencha a coluna 'Fornecedor'"):
        definicao.gravar_linha(base, definicao.linha_vazia(), None)
    linha = {**definicao.linha_vazia(), "nome": "Zeta", "email_para": "z@zeta.com", "idioma": "Espanhol"}
    zeta = definicao.gravar_linha(base, linha, None)
    assert zeta.idioma == Idioma.ES and zeta.ativo  # 'Ativo' vazio = Sim
    with pytest.raises(ErroNegocio, match="idioma inválido"):
        definicao.gravar_linha(base, {**definicao.para_linha(zeta), "idioma": "Alemão"}, zeta)
    with pytest.raises(ErroNegocio, match="Já existe um fornecedor"):
        definicao.gravar_linha(base, {**definicao.linha_vazia(), "nome": "ZETA"}, None)


def test_modelo_controle_cria_rfqs_novas_em_rascunho(base, tmp_path):
    modelo = abas.gerar_modelo_controle(tmp_path / "controle.xlsx")
    livro = load_workbook(modelo)
    aba = livro["Controle"]
    # mesmo pacote para dois fornecedores, sem número de RFQ
    for fornecedor in ("Forn A", "Forn B"):
        aba.append([None, None, "Solicitante X", "Cliente Q", "Projeto Q", fornecedor, "PEÇA 1", "CDC 1", None, "PC"])
        aba.append([None, None, "Solicitante X", "Cliente Q", "Projeto Q", fornecedor, "PEÇA 2", "CDC 2", None, 2])
    livro.save(modelo)

    relatorio = importar_planilha(base, modelo)
    assert relatorio.rfqs_novas == 2 and relatorio.rfqs_importadas == 0 and relatorio.pacotes_criados == 1
    ano = date.today().year
    rfqs = sorted(base.rfqs, key=lambda r: r.numero)
    assert [r.numero for r in rfqs] == [f"RFQ{ano}001", f"RFQ{ano}002"]
    assert all(r.status == StatusRFQ.RASCUNHO and r.prazo is None for r in rfqs)
    assert rfqs[0].pacote_id == rfqs[1].pacote_id
    assert "Enviar e-mails" in "\n".join(relatorio.avisos)
