"""Fase 3: cotação por item, comparativo de propostas e scorecard (dados fictícios)."""

import json
from datetime import date

import pytest
from openpyxl import load_workbook

from rfq_control.armazenamento import BaseDados
from rfq_control.modelos import CotacaoItem, Item, Pacote, StatusRFQ
from rfq_control.servicos import rfq as srv
from rfq_control.servicos.cotacao import comparar, competitividade, definir_cambio
from rfq_control.servicos.exportacao import exportar_comparativo, exportar_rfqs
from rfq_control.servicos.formatos import ler_lead_time, valor_monetario


@pytest.fixture
def pacote(base_populada):
    """Pacote com dois itens (um com volume anual, outro só com quantidade) enviado a f1 e f2."""
    pacote = Pacote(projeto_id="prj1", data=date(2026, 9, 28), itens=[
        Item(ref_op="PECA A", unidade="PC", quantidade=2, volume_anual=1000),
        Item(ref_op="PECA B", unidade="PC", quantidade=3),
    ])
    pacote, rfqs = srv.criar_pacote_com_rfqs(base_populada, pacote, ["f1", "f2"])
    for rfq in rfqs:
        srv.registrar_envio(base_populada, rfq, rfq.idioma, "enviada", date(2026, 9, 1))
    return pacote


def _rfq(base, pacote, fornecedor_id):
    return next(r for r in base.rfqs_do_pacote(pacote.id) if r.fornecedor_id == fornecedor_id)


def _cotar(base, pacote, fornecedor_id, linhas, **campos):
    """Grava a cotação como a tela da RFQ: só o que mudou em relação à RFQ carregada."""
    rfq = _rfq(base, pacote, fornecedor_id)
    a, b = pacote.itens
    cotacao = [CotacaoItem(item_id=item.id, **valores) for item, valores in zip((a, b), linhas) if valores]
    editada = rfq.model_copy(update={"cotacao": cotacao, **campos})
    return srv.salvar_rfq(base, editada, "Dados da RFQ atualizados", carregada=rfq)


# ---------------------------------------------------------------- itens com id


def test_itens_de_pacote_antigo_recebem_ids_estaveis(base, pasta_dados):
    registro = {"id": "p1", "titulo": "Antigo", "projeto_id": "x", "data": "2026-01-02",
                "itens": [{"ref_op": "A"}, {"ref_op": "B"}]}
    (pasta_dados / "pacotes.json").write_text(json.dumps({"registros": [registro]}), encoding="utf-8")
    assert [i.id for i in BaseDados.abrir(pasta_dados).pacotes.obter("p1").itens] == ["i1", "i2"]
    reaberta = BaseDados.abrir(pasta_dados)  # mesmo id a cada leitura, antes de gravar
    pacote = reaberta.pacotes.obter("p1")
    assert [i.id for i in pacote.itens] == ["i1", "i2"]
    # item novo num pacote que já tem ids: id novo, sem reaproveitar o de um item apagado
    editado = Pacote.model_validate({**pacote.model_dump(), "itens": [pacote.itens[1].model_dump(), {"ref_op": "C"}]})
    assert editado.itens[0].id == "i2" and editado.itens[1].id not in ("", "i1", "i2")


# ---------------------------------------------------------------- gravação da cotação


def test_primeira_cotacao_marca_a_rfq_como_respondida(base_populada, pacote):
    salva = _cotar(base_populada, pacote, "f1", [{"preco": 10, "ferramental": 5000, "lead_time_dias": 30}, {}],
                   moeda="BRL", validade=date(2026, 12, 31))
    assert salva.status == StatusRFQ.RESPONDIDA and salva.data_resposta == date.today()
    assert [c.preco for c in salva.cotacao] == [10]  # linha vazia não é gravada
    assert salva.validade == date(2026, 12, 31)
    historico = [e.descricao for e in salva.historico]
    assert "Cotação por item registrada: 1 item(ns) com preço" in historico
    assert "Status alterado de Enviada para Respondida (cotação registrada)" in historico


def test_cotacao_respeita_status_escolhido_e_rascunho(base_populada, pacote, pacote_exemplo):
    salva = _cotar(base_populada, pacote, "f2", [{"preco": 9}, {}], status=StatusRFQ.DECLINADA)
    assert salva.status == StatusRFQ.DECLINADA  # o usuário mudou o status: vale o dele
    _, [rascunho] = srv.criar_pacote_com_rfqs(base_populada, pacote_exemplo, ["f1"])
    editada = rascunho.model_copy(update={"cotacao": [CotacaoItem(item_id=rascunho.id, preco=1)]})
    salva = srv.salvar_rfq(base_populada, editada, carregada=rascunho)
    assert salva.status == StatusRFQ.RASCUNHO and salva.cotacao == []  # item que não é do pacote: ignorado


def test_cotacao_rejeita_valor_negativo():
    with pytest.raises(ValueError):
        CotacaoItem(item_id="i1", preco=-1)
    assert CotacaoItem(item_id="i1", preco=10, frete=0.5, embalagem=0.25).custo_unitario() == 10.75
    assert CotacaoItem(item_id="i1", ferramental=100).custo_unitario() is None


# ---------------------------------------------------------------- comparativo


def test_comparativo_equaliza_e_aponta_os_melhores(base_populada, pacote):
    a, b = pacote.itens
    _cotar(base_populada, pacote, "f1", [{"preco": 10, "frete": 1, "ferramental": 5000, "lead_time_dias": 30},
                                         {"preco": 4, "lead_time_dias": 20}])
    _cotar(base_populada, pacote, "f2", [{"preco": 2, "embalagem": 0.5, "ferramental": 1200, "lead_time_dias": 45},
                                         {"preco": 1, "lead_time_dias": 10}], moeda="USD")
    comparativo = comparar(base_populada, pacote.id)
    um, dois = (_rfq(base_populada, pacote, f) for f in ("f1", "f2"))
    # sem o câmbio do dólar a proposta de f2 fica fora dos destaques
    assert comparativo.moeda == "BRL" and comparativo.moedas_sem_cambio == ["USD"]
    assert comparativo.melhor_total == set() and not comparativo.melhor(a.id, "custo", um.id)

    definir_cambio(base_populada, pacote.id, "brl", {"usd": 5.0, "BRL": 1, "EUR": 0})
    assert base_populada.pacotes.obter(pacote.id).cambios == {"USD": 5.0}
    comparativo = comparar(base_populada, pacote.id)
    proposta_um, proposta_dois = sorted(comparativo.propostas, key=lambda p: p.fornecedor)
    # custo das peças = (preço + frete + embalagem) × volume anual (A) ou quantidade (B)
    assert proposta_um.total_pecas == 11 * 1000 + 4 * 3
    assert proposta_dois.total_pecas == (2.5 * 1000 + 1 * 3) * 5  # convertido de USD
    assert proposta_dois.ferramental == 6000 and proposta_um.lead_time == 30 and proposta_dois.lead_time == 45
    assert comparativo.melhor_total == {um.id}
    assert comparativo.melhor_ferramental == {um.id}
    assert comparativo.melhor_lead_time == {um.id}
    assert comparativo.melhor(b.id, "custo", um.id)  # 4 BRL < 1 USD = 5 BRL
    assert comparativo.melhor(b.id, "lead_time", dois.id)


def test_proposta_incompleta_nao_disputa_o_total(base_populada, pacote):
    a, _ = pacote.itens
    _cotar(base_populada, pacote, "f1", [{"preco": 10}, {"preco": 4}])
    _cotar(base_populada, pacote, "f2", [{"preco": 1}, {}])
    comparativo = comparar(base_populada, pacote.id)
    dois = _rfq(base_populada, pacote, "f2")
    assert {p.rfq.fornecedor_id: p.completa for p in comparativo.propostas} == {"f1": True, "f2": False}
    assert comparativo.melhor_total == set()  # só uma proposta completa: nada a comparar
    assert comparativo.melhor(a.id, "custo", dois.id)  # mas o item cotado é comparado


def test_comparativo_ignora_canceladas_e_lista_as_sem_cotacao(base_populada, pacote):
    _cotar(base_populada, pacote, "f1", [{"preco": 10}, {"preco": 4}])
    comparativo = comparar(base_populada, pacote.id)
    assert [r.fornecedor_id for r in comparativo.sem_cotacao] == ["f2"]
    srv.alterar_status(base_populada, [_rfq(base_populada, pacote, "f2").id], StatusRFQ.CANCELADA)
    comparativo = comparar(base_populada, pacote.id)
    assert comparativo.sem_cotacao == [] and len(comparativo.propostas) == 1
    with pytest.raises(srv.ErroNegocio):
        definir_cambio(base_populada, pacote.id, " ", {})


# ---------------------------------------------------------------- scorecard


def test_scorecard_com_pontualidade_e_menor_preco(base_populada, pacote):
    _cotar(base_populada, pacote, "f1", [{"preco": 10}, {"preco": 4}])
    _cotar(base_populada, pacote, "f2", [{"preco": 20}, {"preco": 4}])
    assert competitividade(base_populada) == {"f1": (1, 1), "f2": (0, 1)}
    # f2 respondeu depois do prazo
    dois = _rfq(base_populada, pacote, "f2")
    base_populada.rfqs.salvar(dois.model_copy(update={"data_resposta": date(2026, 12, 1)}))
    um = _rfq(base_populada, pacote, "f1")
    base_populada.rfqs.salvar(um.model_copy(update={"data_resposta": um.prazo}))
    indicadores = srv.calcular_indicadores(base_populada, hoje=date(2026, 12, 2))
    por_nome = {f.nome: f for f in indicadores.fornecedores}
    assert por_nome["Fornecedor Um"].pontualidade == 1.0 and por_nome["Supplier Two"].pontualidade == 0.0
    assert por_nome["Fornecedor Um"].competitividade == 1.0 and por_nome["Supplier Two"].competitividade == 0.0


# ---------------------------------------------------------------- Excel e formatos


def test_exportacoes_trazem_a_cotacao(base_populada, pacote, tmp_path):
    _cotar(base_populada, pacote, "f1", [{"preco": 10, "ferramental": 500, "lead_time_dias": 30}, {"preco": 4}],
           validade=date(2026, 12, 31))
    _cotar(base_populada, pacote, "f2", [{"preco": 12}, {"preco": 5}])
    livro = load_workbook(exportar_rfqs(base_populada, tmp_path / "rfqs.xlsx"))
    cabecalho = [c.value for c in livro["Itens"][1]]
    assert cabecalho[11:17] == ["Moeda", "Preço unitário", "Ferramental", "Lead time (dias)", "Frete", "Embalagem"]
    linhas = [r for r in livro["Itens"].iter_rows(min_row=2, values_only=True) if r[0] == _rfq(base_populada, pacote, "f1").numero]
    assert linhas[0][11:15] == ("BRL", 10, 500, 30)
    assert "Validade da proposta" in [c.value for c in livro["RFQs"][1]]

    comparativo = load_workbook(exportar_comparativo(base_populada, pacote.id, tmp_path / "comparativo.xlsx"))
    assert comparativo.sheetnames == ["Resumo", "Itens"]
    resumo = list(comparativo["Resumo"].iter_rows(min_row=2, max_row=3, values_only=True))
    assert {r[1]: r[7] for r in resumo} == {"Fornecedor Um": 10 * 1000 + 4 * 3, "Supplier Two": 12 * 1000 + 5 * 3}
    melhor = next(l for l in range(2, 4) if comparativo["Resumo"][f"B{l}"].value == "Fornecedor Um")
    assert comparativo["Resumo"][f"H{melhor}"].fill.fgColor.rgb.endswith("C6EFCE")


def test_formatos_de_preco_e_lead_time():
    assert [valor_monetario(v) for v in (0.1234, 1234.5, 2, None)] == ["0,1234", "1.234,50", "2,00", ""]
    assert [ler_lead_time(t) for t in ("30", "30 dias", "6 semanas", "6 sem.", "8 weeks", "2 meses", "")] == [
        30, 30, 42, 42, 56, 60, None]
    with pytest.raises(ValueError):
        ler_lead_time("logo")
