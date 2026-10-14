"""Abas da base de dados (as antigas abas da planilha) e cadastro em massa.

Cada aba de cadastro é descrita por uma ``DefinicaoAba``: as colunas que aparecem na grade e nas
planilhas de importação/exportação, e como converter uma linha (textos) em registro e vice-versa.
A mesma definição é usada pela tela "Base de dados" (edição direto nas células) e pela
importação em massa de arquivos Excel/CSV.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from collections.abc import Callable, Iterator
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel, ValidationError

from ..armazenamento import BaseDados
from ..modelos import Commodity, Contato, Feriado, Fornecedor, Idioma, Projeto, Solicitante, agora, separar_emails
from . import cadastros
from .formatos import data_curta, ler_data, nomes_de_opcoes, normalizar_nome
from .leitor_xlsx import ErroPlanilha, LeitorXlsx, Valor, data_excel, letras_coluna
from .rfq import ErroNegocio

SIM, NAO = "Sim", "Não"
IDIOMAS = tuple(i.rotulo for i in Idioma)  # ("Português", "Espanhol", "Inglês")


# ---------------------------------------------------------------- colunas e conversões


@dataclass(frozen=True)
class ColunaAba:
    chave: str
    titulo: str
    tipo: str = "texto"  # texto | data | inteiro | lista | multilista | booleano | emails
    opcoes: tuple[str, ...] = ()
    obrigatoria: bool = False
    largura: int = 150
    apelidos: tuple[str, ...] = ()  # outros cabeçalhos aceitos na importação (inclui os da planilha antiga)
    dica: str = ""
    somente_leitura: bool = False  # calculada a partir de outros cadastros (exibida, nunca gravada)
    opcoes_base: Callable[[BaseDados], list[str]] | None = None  # opções que vêm de outro cadastro
    criar_opcao: Callable[[BaseDados, str], Any] | None = None  # importação: cria a opção que faltar

    def opcoes_de(self, base: BaseDados | None) -> tuple[str, ...]:
        if self.opcoes_base and base is not None:
            return tuple(self.opcoes_base(base))
        return self.opcoes


def chave_cabecalho(texto: Any) -> str:
    """'  EMAIL (To) ' → 'emailto'; 'Descrição' → 'descricao'."""
    sem_acento = "".join(
        c for c in unicodedata.normalize("NFD", str(texto or "").casefold()) if unicodedata.category(c) != "Mn"
    )
    return re.sub(r"[^a-z0-9]", "", sem_acento)


def ler_booleano(texto: str, padrao: bool = True) -> bool:
    chave = chave_cabecalho(texto)
    if not chave:
        return padrao
    if chave in {"sim", "s", "yes", "y", "true", "verdadeiro", "1", "x", "ativo"}:
        return True
    if chave in {"nao", "n", "no", "false", "falso", "0", "inativo"}:
        return False
    raise ValueError(f"valor inválido: '{texto}' (use Sim ou Não)")


def texto_booleano(valor: bool) -> str:
    return SIM if valor else NAO


def _limpar_responsavel(nome: str) -> str:
    # A planilha antiga preenchia RESPONSÁVEL com "fornecedor" quando não havia nome.
    return "" if normalizar_nome(nome) in {"fornecedor", "supplier", "proveedor"} else nome.strip()


def _validar(modelo: type[BaseModel], dados: dict, nomes: dict[str, str]) -> Any:
    try:
        return modelo.model_validate(dados)
    except ValidationError as erro:
        partes = []
        for detalhe in erro.errors()[:3]:
            campo = str(detalhe["loc"][0]) if detalhe["loc"] else ""
            mensagem = str(detalhe["msg"]).replace("Value error, ", "")
            partes.append(f"{nomes.get(campo, campo)}: {mensagem}" if campo else mensagem)
        raise ErroNegocio("; ".join(partes)) from erro


@dataclass
class DefinicaoAba:
    chave: str
    titulo: str
    nome_legado: str
    descricao: str
    colunas: list[ColunaAba]
    listar: Callable[[BaseDados], list]
    para_linha: Callable[[BaseDados, Any], dict[str, str]]
    de_linha: Callable[[BaseDados, dict[str, str], Any | None], Any]
    salvar: Callable[[BaseDados, Any], Any]
    chave_linha: Callable[[dict[str, str]], str]
    chave_registro: Callable[[Any], str]
    coluna_a_sem_cabecalho: str | None = None  # ex.: aba Feriados antiga (só datas na coluna A)

    def coluna(self, chave: str) -> ColunaAba:
        return next(c for c in self.colunas if c.chave == chave)

    @property
    def obrigatoria(self) -> ColunaAba:
        return next(c for c in self.colunas if c.obrigatoria)

    @property
    def editaveis(self) -> list[ColunaAba]:
        return [c for c in self.colunas if not c.somente_leitura]

    def nomes_campos(self) -> dict[str, str]:
        nomes = {c.chave: c.titulo for c in self.colunas}
        nomes.update({"contatos": "E-mail", "data": nomes.get("data", "Data"), "commodity_ids": "Commodity"})
        return nomes

    def linha_vazia(self) -> dict[str, str]:
        return {c.chave: "" for c in self.colunas}

    def registro_de_linha(self, base: BaseDados, linha: dict[str, str], existente: Any | None) -> Any:
        """Converte e valida (sem gravar). Colunas ausentes em ``linha`` mantêm o valor atual."""
        somente_leitura = {c.chave for c in self.colunas if c.somente_leitura}
        linha = {chave: valor for chave, valor in linha.items() if chave not in somente_leitura}
        obrigatoria = self.obrigatoria
        if obrigatoria.chave in linha and not str(linha[obrigatoria.chave]).strip():
            raise ErroNegocio(f"Preencha a coluna '{obrigatoria.titulo}'.")
        if existente is None and not str(linha.get(obrigatoria.chave, "")).strip():
            raise ErroNegocio(f"Preencha a coluna '{obrigatoria.titulo}'.")
        try:
            return self.de_linha(base, linha, existente)
        except ErroNegocio:
            raise
        except ValueError as erro:  # idioma, Sim/Não…
            raise ErroNegocio(str(erro)) from erro

    def gravar_linha(self, base: BaseDados, linha: dict[str, str], existente: Any | None) -> Any:
        return self.salvar(base, self.registro_de_linha(base, linha, existente))


# ---------------------------------------------------------------- fornecedores


def _linha_fornecedor(base: BaseDados, f: Fornecedor) -> dict[str, str]:
    commodities = base.commodities_do_fornecedor(f)
    return {
        "nome": f.nome,
        "email_para": "; ".join(f.emails_para()),
        "email_cc": "; ".join(f.emails_copia()),
        "commodities": "; ".join(c.nome for c in commodities),
        "email_buyer": "; ".join(dict.fromkeys(c.email for c in commodities if c.email)),
        "idioma": f.idioma.rotulo,
    }


def _contatos_da_linha(linha: dict[str, str], existente: Fornecedor | None) -> list[Contato]:
    """Contatos a partir das colunas de e-mail; nomes e telefones já cadastrados são mantidos."""
    atual_para = existente.emails_para() if existente else []
    atual_cc = existente.emails_copia() if existente else []
    para = separar_emails(linha["email_para"]) if "email_para" in linha else atual_para
    cc = separar_emails(linha["email_cc"]) if "email_cc" in linha else atual_cc
    cc = [e for e in dict.fromkeys(cc) if e not in para]
    para = list(dict.fromkeys(para))
    if existente and para == atual_para and cc == atual_cc:
        return existente.contatos  # nada mudou: preserva tudo
    antigos = {c.email: c for c in existente.contatos if c.email} if existente else {}
    principal = next((c for c in existente.contatos if c.email and not c.copia), None) if existente else None
    contatos: list[Contato] = []
    for posicao, email in enumerate(para):
        antigo = antigos.get(email)
        if antigo is None and posicao == 0 and principal and principal.email not in para:
            antigo = principal  # trocou o e-mail do contato principal: mantém o nome usado na saudação
        contatos.append(_contato(antigo.nome if antigo else "", email, False, antigo.telefone if antigo else ""))
    for email in cc:
        antigo = antigos.get(email)
        contatos.append(_contato(antigo.nome if antigo else "", email, True, antigo.telefone if antigo else ""))
    sem_email = [c for c in existente.contatos if not c.email] if existente else []
    return contatos + sem_email


def _contato(nome: str, email: str, copia: bool, telefone: str) -> Contato:
    try:
        return Contato(nome=nome, email=email, copia=copia, telefone=telefone)
    except ValidationError as erro:
        raise ErroNegocio(f"E-mail inválido: {email}") from erro


def _ids_commodities(base: BaseDados, texto: str) -> list[str]:
    por_nome = {normalizar_nome(c.nome): c for c in base.commodities}
    ids, faltando = [], []
    for nome in nomes_de_opcoes(texto, set(por_nome)):
        commodity = por_nome.get(normalizar_nome(nome))
        if commodity is None:
            faltando.append(nome)
        elif commodity.id not in ids:
            ids.append(commodity.id)
    if faltando:
        raise ErroNegocio(
            "Commodity não cadastrada: " + ", ".join(faltando) + ". Cadastre na aba Commodities."
        )
    return ids


def _fornecedor_de_linha(base: BaseDados, linha: dict[str, str], existente: Fornecedor | None) -> Fornecedor:
    dados = existente.model_dump() if existente else {}
    if "nome" in linha:
        dados["nome"] = linha["nome"]
    if str(linha.get("idioma", "")).strip():
        dados["idioma"] = Idioma.de_texto(linha["idioma"])
    if "commodities" in linha:
        dados["commodity_ids"] = _ids_commodities(base, linha["commodities"])
    dados["contatos"] = [c.model_dump() for c in _contatos_da_linha(linha, existente)]
    return _validar(Fornecedor, dados, ABA_FORNECEDORES.nomes_campos())


def _nomes_commodities(base: BaseDados) -> list[str]:
    return sorted((c.nome for c in base.commodities), key=str.casefold)


def _criar_commodity(base: BaseDados, nome: str) -> Commodity:
    return cadastros.salvar_commodity(base, Commodity(nome=nome))


ABA_FORNECEDORES = DefinicaoAba(
    chave="fornecedores",
    titulo="Fornecedores",
    nome_legado="CadastroContatos",
    descricao="Fornecedores, e-mails, commodities e idioma dos e-mails (antiga aba CadastroContatos).",
    colunas=[
        ColunaAba("nome", "Fornecedor", obrigatoria=True, largura=220,
                  apelidos=("fornecedor", "supplier", "supplierdistributor", "nome", "razaosocial")),
        ColunaAba("email_para", "E-mail (Para)", "emails", largura=230,
                  apelidos=("emailto", "emailpara", "email", "to", "para", "destinatario"),
                  dica="Quem recebe o e-mail com o escopo. Vários e-mails separados por ;"),
        ColunaAba("email_cc", "E-mails Supplier (Cc)", "emails", largura=230,
                  apelidos=("emailsuppliercc", "suppliercc", "emailcc", "cc", "copia", "emailcopia", "buyeremail"),
                  dica="Pessoas do fornecedor em cópia. Vários e-mails separados por ;"),
        ColunaAba("commodities", "Commodity", "multilista", largura=210,
                  apelidos=("commodity", "commodities", "categoria", "category", "segmento"),
                  dica="Uma ou mais commodities do fornecedor (cadastradas na aba Commodities), separadas por ;",
                  opcoes_base=_nomes_commodities, criar_opcao=_criar_commodity),
        ColunaAba("email_buyer", "E-mail Commodity Buyer (Cc)", "emails", largura=230, somente_leitura=True,
                  apelidos=("emailcommoditybuyer", "commoditybuyer", "commoditybuyercc"),
                  dica="Vem da aba Commodities: o Commodity Buyer de cada commodity vai em cópia nos e-mails. "
                       "Para mudar, edite a commodity."),
        ColunaAba("idioma", "Idioma", "lista", IDIOMAS, largura=110, apelidos=("lingua", "language", "idiomas"),
                  dica="Idioma dos e-mails de RFQ para este fornecedor"),
    ],
    listar=lambda base: sorted(base.fornecedores, key=lambda f: f.nome.casefold()),
    para_linha=_linha_fornecedor,
    de_linha=_fornecedor_de_linha,
    salvar=cadastros.salvar_fornecedor,
    chave_linha=lambda linha: normalizar_nome(linha.get("nome", "")),
    chave_registro=lambda f: normalizar_nome(f.nome),
)


# ---------------------------------------------------------------- commodities


def _commodity_de_linha(base: BaseDados, linha: dict[str, str], existente: Commodity | None) -> Commodity:
    dados = existente.model_dump() if existente else {}
    for chave in ("nome", "comprador", "email"):
        if chave in linha:
            dados[chave] = linha[chave]
    return _validar(Commodity, dados, ABA_COMMODITIES.nomes_campos())


ABA_COMMODITIES = DefinicaoAba(
    chave="commodities",
    titulo="Commodities",
    nome_legado="Commodity",
    descricao="Commodities e o Commodity Buyer de cada uma: ele vai em cópia (Cc) nos e-mails dos pacotes e "
              "dos fornecedores da commodity.",
    colunas=[
        ColunaAba("nome", "Commodity", obrigatoria=True, largura=240,
                  apelidos=("commodity", "nome", "categoria", "category")),
        ColunaAba("comprador", "Commodity Buyer", largura=220,
                  apelidos=("commoditybuyer", "buyer", "comprador", "responsavel")),
        ColunaAba("email", "E-mail do Commodity Buyer", largura=260,
                  apelidos=("emailcommoditybuyer", "emailbuyer", "buyeremail", "email", "emailcomprador")),
        ColunaAba("fornecedores", "Fornecedores", largura=320, somente_leitura=True,
                  dica="Fornecedores com esta commodity (escolha na aba Fornecedores)"),
    ],
    listar=lambda base: sorted(base.commodities, key=lambda c: c.nome.casefold()),
    para_linha=lambda base, c: {
        "nome": c.nome, "comprador": c.comprador, "email": c.email,
        "fornecedores": ", ".join(sorted((f.nome for f in base.fornecedores_da_commodity(c.id)), key=str.casefold)),
    },
    de_linha=_commodity_de_linha,
    salvar=cadastros.salvar_commodity,
    chave_linha=lambda linha: normalizar_nome(linha.get("nome", "")),
    chave_registro=lambda c: normalizar_nome(c.nome),
)


# ---------------------------------------------------------------- projetos


def _inteiro(texto: str, titulo: str) -> int | None:
    texto = str(texto or "").strip()
    if not texto:
        return None
    try:
        return int(float(texto.replace(",", ".")))
    except ValueError as erro:
        raise ErroNegocio(f"{titulo}: '{texto}' não é um número inteiro.") from erro


def _data(texto: str, titulo: str) -> date | None:
    try:
        return ler_data(str(texto or ""))
    except ValueError as erro:
        raise ErroNegocio(f"{titulo}: {erro}") from erro


def _projeto_de_linha(base: BaseDados, linha: dict[str, str], existente: Projeto | None) -> Projeto:
    dados = existente.model_dump() if existente else {}
    for chave in ("nome", "cliente", "planta"):
        if chave in linha:
            dados[chave] = linha[chave]
    if "sop" in linha:
        dados["sop"] = _data(linha["sop"], "SOP")
    if "lifetime_anos" in linha:
        dados["lifetime_anos"] = _inteiro(linha["lifetime_anos"], "Lifetime (anos)")
    if "ativo" in linha:
        dados["ativo"] = ler_booleano(linha["ativo"])
    return _validar(Projeto, dados, ABA_PROJETOS.nomes_campos())


ABA_PROJETOS = DefinicaoAba(
    chave="projetos",
    titulo="Projetos",
    nome_legado="CadastroProjetos",
    descricao="Projetos, clientes e plantas (antiga aba CadastroProjetos).",
    colunas=[
        ColunaAba("planta", "Planta", largura=180, apelidos=("plantaop", "planta", "plant")),
        ColunaAba("cliente", "Cliente", largura=140, apelidos=("cliente", "customer")),
        ColunaAba("nome", "Projeto", obrigatoria=True, largura=240, apelidos=("projeto", "project", "nome")),
        ColunaAba("sop", "SOP", "data", largura=110, apelidos=("sop", "startofproduction"), dica="dd/mm/aaaa"),
        ColunaAba("lifetime_anos", "Lifetime (anos)", "inteiro", largura=120,
                  apelidos=("lifetimeyears", "lifetime", "lifetimeanos", "vidautil")),
        ColunaAba("ativo", "Ativo", "booleano", (SIM, NAO), largura=70, apelidos=("ativo", "active")),
    ],
    listar=lambda base: sorted(base.projetos, key=lambda p: p.nome.casefold()),
    para_linha=lambda base, p: {
        "planta": p.planta, "cliente": p.cliente, "nome": p.nome, "sop": data_curta(p.sop),
        "lifetime_anos": str(p.lifetime_anos or ""), "ativo": texto_booleano(p.ativo),
    },
    de_linha=_projeto_de_linha,
    salvar=cadastros.salvar_projeto,
    chave_linha=lambda linha: normalizar_nome(linha.get("nome", "")),
    chave_registro=lambda p: normalizar_nome(p.nome),
)


# ---------------------------------------------------------------- solicitantes


def _solicitante_de_linha(base: BaseDados, linha: dict[str, str], existente: Solicitante | None) -> Solicitante:
    dados = existente.model_dump() if existente else {}
    for chave in ("nome", "email"):
        if chave in linha:
            dados[chave] = linha[chave]
    if "ativo" in linha:
        dados["ativo"] = ler_booleano(linha["ativo"])
    return _validar(Solicitante, dados, ABA_SOLICITANTES.nomes_campos())


ABA_SOLICITANTES = DefinicaoAba(
    chave="solicitantes",
    titulo="Solicitantes",
    nome_legado="Requester",
    descricao="Pessoas que solicitam as cotações.",
    colunas=[
        ColunaAba("nome", "Nome", obrigatoria=True, largura=240,
                  apelidos=("nome", "solicitante", "requester", "comprador", "buyer")),
        ColunaAba("email", "E-mail", largura=260, apelidos=("email", "mail")),
        ColunaAba("ativo", "Ativo", "booleano", (SIM, NAO), largura=70, apelidos=("ativo", "active")),
    ],
    listar=lambda base: sorted(base.solicitantes, key=lambda s: s.nome.casefold()),
    para_linha=lambda base, s: {"nome": s.nome, "email": s.email, "ativo": texto_booleano(s.ativo)},
    de_linha=_solicitante_de_linha,
    salvar=cadastros.salvar_solicitante,
    chave_linha=lambda linha: normalizar_nome(linha.get("nome", "")),
    chave_registro=lambda s: normalizar_nome(s.nome),
)


# ---------------------------------------------------------------- feriados


def _feriado_de_linha(base: BaseDados, linha: dict[str, str], existente: Feriado | None) -> Feriado:
    dados = existente.model_dump() if existente else {}
    if "data" in linha:
        dados["data"] = _data(linha["data"], "Data")
    if "descricao" in linha:
        dados["descricao"] = linha["descricao"]
    return _validar(Feriado, dados, ABA_FERIADOS.nomes_campos())


def _chave_data(texto: str) -> str:
    try:
        dia = ler_data(str(texto or ""))
    except ValueError:
        return ""
    return dia.isoformat() if dia else ""


ABA_FERIADOS = DefinicaoAba(
    chave="feriados",
    titulo="Feriados",
    nome_legado="Feriados",
    descricao="Datas desconsideradas no cálculo do prazo em dias úteis.",
    colunas=[
        ColunaAba("data", "Data", "data", obrigatoria=True, largura=120, apelidos=("data", "date", "dia"),
                  dica="dd/mm/aaaa"),
        ColunaAba("descricao", "Descrição", largura=320, apelidos=("descricao", "feriado", "nome", "holiday")),
    ],
    listar=lambda base: sorted(base.feriados, key=lambda f: f.data, reverse=True),
    para_linha=lambda base, f: {"data": data_curta(f.data), "descricao": f.descricao},
    de_linha=_feriado_de_linha,
    salvar=cadastros.salvar_feriado,
    chave_linha=lambda linha: _chave_data(linha.get("data", "")),
    chave_registro=lambda f: f.data.isoformat(),
    coluna_a_sem_cabecalho="data",
)

ABAS_CADASTRO = [ABA_FORNECEDORES, ABA_COMMODITIES, ABA_PROJETOS, ABA_SOLICITANTES, ABA_FERIADOS]


# ---------------------------------------------------------------- leitura de arquivos


def _texto_celula(valor: Valor, coluna: ColunaAba) -> str:
    if valor is None:
        return ""
    if coluna.tipo == "data" and isinstance(valor, float):
        dia = data_excel(valor)
        return data_curta(dia) if dia else ""
    if coluna.tipo == "booleano" and isinstance(valor, bool):
        return texto_booleano(valor)
    if isinstance(valor, float):
        return str(int(valor)) if valor.is_integer() else str(valor)
    return str(valor).strip()


def _linhas_csv(caminho: Path) -> list[list[str]]:
    bruto = caminho.read_bytes()
    for codificacao in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = bruto.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    primeira = texto.splitlines()[0] if texto.strip() else ""
    separador = max((";", ",", "\t"), key=primeira.count) if primeira else ";"
    return [linha for linha in csv.reader(io.StringIO(texto), delimiter=separador)]


def _mapear_cabecalho(definicao: DefinicaoAba, cabecalho: dict[str, Valor]) -> dict[str, list[str]]:
    """{chave da coluna: [colunas do arquivo]} — uma chave pode juntar mais de uma coluna (ex.: Cc + Buyer)."""
    mapa: dict[str, list[str]] = {}
    for letra, texto in cabecalho.items():
        chave_arquivo = chave_cabecalho(texto)
        if not chave_arquivo:
            continue
        for coluna in definicao.colunas:
            if chave_arquivo in {chave_cabecalho(coluna.titulo), *coluna.apelidos}:
                mapa.setdefault(coluna.chave, []).append(letra)
                break
    return mapa


@dataclass
class _Tabela:
    linhas: list[tuple[int, dict[str, Valor]]]  # (número da linha no arquivo, {coluna: valor})
    mapa: dict[str, list[str]]
    cabecalhos: dict[str, str]


def _localizar_tabela(
    definicao: DefinicaoAba, linhas: Iterator[tuple[int, dict[str, Valor]]], aceitar_sem_cabecalho: bool
) -> _Tabela | None:
    """Procura o cabeçalho nas 10 primeiras linhas; sem cabeçalho, usa a coluna A se a aba permitir."""
    iniciais = []
    for numero, valores in linhas:
        iniciais.append((numero, valores))
        mapa = _mapear_cabecalho(definicao, valores)
        if definicao.obrigatoria.chave in mapa:
            cabecalhos = {letra: str(texto) for letra, texto in valores.items()}
            return _Tabela(list(linhas), mapa, cabecalhos)
        if len(iniciais) >= 10:
            break
    if aceitar_sem_cabecalho and definicao.coluna_a_sem_cabecalho and iniciais:
        return _Tabela(iniciais + list(linhas), {definicao.coluna_a_sem_cabecalho: ["A"]}, {"A": "(coluna A)"})
    return None


def _linhas_do_arquivo(definicao: DefinicaoAba, caminho: Path) -> _Tabela:
    sufixo = caminho.suffix.lower()
    if sufixo in (".csv", ".txt"):
        linhas = (
            (numero, {letras_coluna(i): v for i, v in enumerate(linha, start=1) if v.strip()})
            for numero, linha in enumerate(_linhas_csv(caminho), start=1)
        )
        tabela = _localizar_tabela(definicao, (l for l in linhas if l[1]), aceitar_sem_cabecalho=True)
        if tabela is None:
            raise ErroNegocio(_mensagem_sem_cabecalho(definicao))
        return tabela

    if sufixo not in (".xlsx", ".xlsm"):
        raise ErroNegocio("Formato não suportado. Use planilhas .xlsx/.xlsm ou arquivos .csv.")
    with LeitorXlsx(caminho) as leitor:
        preferidos = {chave_cabecalho(n) for n in (definicao.titulo, definicao.chave, definicao.nome_legado)}
        # Primeiro a aba com o nome do cadastro (ex.: 'CadastroContatos' na planilha antiga), depois as demais.
        for aba in sorted(leitor.abas, key=lambda a: chave_cabecalho(a) not in preferidos):
            # Ler sem cabeçalho só faz sentido na aba de mesmo nome ou em arquivo de uma aba só.
            sem_cabecalho = chave_cabecalho(aba) in preferidos or len(leitor.abas) == 1
            tabela = _localizar_tabela(definicao, leitor.linhas(aba, parar_apos_vazias=1000), sem_cabecalho)
            if tabela:
                return tabela
    raise ErroNegocio(_mensagem_sem_cabecalho(definicao))


def _mensagem_sem_cabecalho(definicao: DefinicaoAba) -> str:
    return (
        f"Não encontrei a coluna '{definicao.obrigatoria.titulo}' no arquivo. Use o botão 'Baixar modelo' "
        "para ver o formato esperado."
    )


# ---------------------------------------------------------------- importação em massa


@dataclass
class ResultadoImportacaoAba:
    criados: int = 0
    atualizados: int = 0
    sem_alteracao: int = 0
    opcoes_criadas: list[str] = field(default_factory=list)  # ex.: commodities novas criadas na importação
    erros: list[str] = field(default_factory=list)
    colunas_reconhecidas: list[str] = field(default_factory=list)
    colunas_ignoradas: list[str] = field(default_factory=list)
    pasta_backup: Path | None = None

    def resumo(self) -> str:
        partes = [
            f"Novos: {self.criados}",
            f"Atualizados: {self.atualizados}",
            f"Sem alteração: {self.sem_alteracao}",
        ]
        if self.opcoes_criadas:
            partes.append("Commodities criadas (complete o Commodity Buyer na aba Commodities): "
                          + ", ".join(self.opcoes_criadas))
        if self.erros:
            partes.append(f"Linhas com erro: {len(self.erros)}")
        return "\n".join(partes)


def _criar_opcoes_faltantes(base: BaseDados, definicao: DefinicaoAba, tabela: _Tabela) -> list[str]:
    """Importação em massa: cria as opções citadas no arquivo que ainda não existem (ex.: commodities)."""
    criadas: list[str] = []
    for chave, letras in tabela.mapa.items():
        coluna = definicao.coluna(chave)
        if coluna.criar_opcao is None:
            continue
        existentes = {normalizar_nome(o) for o in coluna.opcoes_de(base)}
        for _numero, valores in tabela.linhas:
            for letra in letras:
                for nome in nomes_de_opcoes(_texto_celula(valores.get(letra), coluna), existentes):
                    if normalizar_nome(nome) not in existentes:
                        coluna.criar_opcao(base, nome)
                        existentes.add(normalizar_nome(nome))
                        criadas.append(nome)
    return criadas


def _sem_id(registro: BaseModel) -> dict:
    dados = registro.model_dump(mode="json")
    dados.pop("id", None)
    return dados


def importar_arquivo(base: BaseDados, definicao: DefinicaoAba, caminho: Path | str) -> ResultadoImportacaoAba:
    """Cadastro em massa: cria ou atualiza (pelo nome/data) os registros de uma planilha ou CSV."""
    caminho = Path(caminho)
    try:
        tabela = _linhas_do_arquivo(definicao, caminho)
    except ErroPlanilha as erro:
        raise ErroNegocio(str(erro)) from erro
    resultado = ResultadoImportacaoAba()
    usadas = {letra for letras in tabela.mapa.values() for letra in letras}
    resultado.colunas_ignoradas = [t for l, t in tabela.cabecalhos.items() if l not in usadas and t.strip()]
    # colunas calculadas (ex.: e-mail do Commodity Buyer) são reconhecidas, mas não importadas
    tabela.mapa = {chave: letras for chave, letras in tabela.mapa.items() if not definicao.coluna(chave).somente_leitura}
    resultado.colunas_reconhecidas = [definicao.coluna(c).titulo for c in tabela.mapa]
    resultado.pasta_backup = base.backup(f"antes_importacao_{definicao.chave}_{agora():%Y-%m-%d_%H%M%S}")
    resultado.opcoes_criadas = _criar_opcoes_faltantes(base, definicao, tabela)

    existentes = {definicao.chave_registro(r): r for r in definicao.listar(base)}
    for numero, valores in tabela.linhas:
        linha: dict[str, str] = {}
        for chave, letras in tabela.mapa.items():
            coluna = definicao.coluna(chave)
            textos = [_texto_celula(valores.get(letra), coluna) for letra in letras]
            separador = "; " if coluna.tipo in ("emails", "multilista") else " "
            linha[chave] = separador.join(t for t in textos if t)
        if not any(linha.values()):
            continue
        chave = definicao.chave_linha(linha)
        if not chave:
            resultado.erros.append(f"Linha {numero}: coluna '{definicao.obrigatoria.titulo}' vazia ou inválida.")
            continue
        existente = existentes.get(chave)
        try:
            registro = definicao.registro_de_linha(base, linha, existente)
            if existente is not None and _sem_id(registro) == _sem_id(existente):
                resultado.sem_alteracao += 1
                continue
            salvo = definicao.salvar(base, registro)
        except (ErroNegocio, ValueError) as erro:
            resultado.erros.append(f"Linha {numero}: {erro}")
            continue
        existentes[chave] = salvo
        if existente is None:
            resultado.criados += 1
        else:
            resultado.atualizados += 1
    return resultado


# ---------------------------------------------------------------- exportação e modelo

_FONTE_CABECALHO = Font(bold=True, color="FFFFFF")
_FUNDO_CABECALHO = PatternFill("solid", fgColor="2E2E2E")
_FUNDO_OBRIGATORIA = PatternFill("solid", fgColor="C0392B")


def _planilha_da_aba(
    definicao: DefinicaoAba, linhas: list[dict[str, str]], destino: Path, instrucoes: bool,
    base: BaseDados | None = None,
) -> Path:
    livro = Workbook()
    aba = livro.active
    aba.title = definicao.titulo[:31]
    # o modelo de importação só tem as colunas editáveis; a exportação mostra também as calculadas
    colunas = definicao.editaveis if instrucoes else definicao.colunas
    aba.append([c.titulo for c in colunas])
    for celula, coluna in zip(aba[1], colunas):
        celula.font = _FONTE_CABECALHO
        celula.fill = _FUNDO_OBRIGATORIA if coluna.obrigatoria else _FUNDO_CABECALHO
    for indice, coluna in enumerate(colunas, start=1):
        aba.column_dimensions[get_column_letter(indice)].width = max(12, coluna.largura // 7)
        if coluna.opcoes:
            validacao = DataValidation(type="list", formula1='"' + ",".join(coluna.opcoes) + '"', allow_blank=True)
            validacao.error = "Escolha um valor da lista: " + ", ".join(coluna.opcoes)
            validacao.prompt = coluna.dica or ("Opções: " + ", ".join(coluna.opcoes))
            aba.add_data_validation(validacao)
            letra = get_column_letter(indice)
            validacao.add(f"{letra}2:{letra}5000")
    for linha in linhas:
        aba.append([linha.get(c.chave, "") for c in colunas])
    aba.freeze_panes = "A2"
    if instrucoes:
        guia = livro.create_sheet("Instruções")
        guia.column_dimensions["A"].width = 28
        guia.column_dimensions["B"].width = 90
        guia.append([f"Importação em massa — {definicao.titulo}"])
        guia["A1"].font = Font(bold=True, size=13)
        guia.append([])
        guia.append(["Como usar", "Preencha uma linha por registro na primeira aba e importe pelo botão "
                                  "'Importar Excel/CSV' na tela Base de dados."])
        guia.append(["Registros existentes", "Linhas com o mesmo nome (ou data, nos feriados) atualizam o "
                                             "registro existente; colunas removidas da planilha não são alteradas."])
        if any(c.tipo == "multilista" for c in colunas):
            guia.append(["Vários valores", "Separe com ; (ex.: Ferramentaria; Dispositivos). Commodities que "
                                           "não existirem são criadas na importação."])
        guia.append(["Coluna em vermelho", f"'{definicao.obrigatoria.titulo}' é obrigatória."])
        for coluna in colunas:
            detalhes = coluna.dica or ""
            opcoes = coluna.opcoes_de(base)
            if opcoes:
                detalhes = (detalhes + " — " if detalhes else "") + "valores: " + ", ".join(opcoes)
            if coluna.tipo == "data":
                detalhes = detalhes or "dd/mm/aaaa"
            guia.append([coluna.titulo, detalhes])
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    livro.save(destino)
    return destino


def gerar_modelo(definicao: DefinicaoAba, destino: Path | str, base: BaseDados | None = None) -> Path:
    return _planilha_da_aba(definicao, [], Path(destino), instrucoes=True, base=base)


def exportar_aba(base: BaseDados, definicao: DefinicaoAba, destino: Path | str) -> Path:
    linhas = [definicao.para_linha(base, r) for r in definicao.listar(base)]
    return _planilha_da_aba(definicao, linhas, Path(destino), instrucoes=False)


CABECALHO_CONTROLE = [
    "RFQ Nº", "DATA", "SOLICITANTE", "CLIENTE", "PROJETO", "FORNECEDOR", "OP - Ref", "CDC - Ref",
    "PART DESCRIPTION", "UM / QTY", "Volume annual",
]


def gerar_modelo_controle(destino: Path | str) -> Path:
    """Planilha no formato da aba Controle antiga, para criar RFQs em massa."""
    livro = Workbook()
    aba = livro.active
    aba.title = "Controle"
    aba.append(CABECALHO_CONTROLE)
    for celula in aba[1]:
        celula.font = _FONTE_CABECALHO
        celula.fill = _FUNDO_CABECALHO
    for indice, largura in enumerate([14, 12, 18, 14, 24, 24, 28, 32, 28, 12, 14], start=1):
        aba.column_dimensions[get_column_letter(indice)].width = largura
    aba.freeze_panes = "A2"
    guia = livro.create_sheet("Instruções")
    guia.column_dimensions["A"].width = 24
    guia.column_dimensions["B"].width = 100
    textos = [
        ("Importação de RFQs", ""),
        ("", ""),
        ("Uma linha por item", "Repita RFQ/fornecedor/projeto em cada linha de item, como na aba Controle antiga."),
        ("RFQ Nº preenchido", "A RFQ é registrada com esse número como já enviada (histórico)."),
        ("RFQ Nº em branco", "O sistema cria RFQs novas em rascunho, numeradas automaticamente — uma por "
                             "fornecedor — prontas para o envio dos e-mails."),
        ("Mesmos itens", "Fornecedores com o mesmo projeto e os mesmos itens ficam no mesmo pacote de cotação."),
        ("UM / QTY", "Unidade (EA, PC…) ou quantidade numérica."),
        ("Cadastros", "Projetos, fornecedores e solicitantes que não existirem são criados automaticamente."),
    ]
    for titulo, texto in textos:
        guia.append([titulo, texto])
    guia["A1"].font = Font(bold=True, size=13)
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    livro.save(destino)
    return destino
