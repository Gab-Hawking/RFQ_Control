"""Exportação das RFQs para Excel (.xlsx)."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..armazenamento import BaseDados
from ..modelos import RFQ
from .cotacao import comparar

_CABECALHO = Font(bold=True, color="FFFFFF")
_FUNDO = PatternFill("solid", fgColor="1F4E79")
_MELHOR = PatternFill("solid", fgColor="C6EFCE")  # melhor valor do comparativo
_FORMATO_VALOR = "#,##0.00##"


def _formatar_aba(aba, larguras: list[int]) -> None:
    for celula in aba[1]:
        celula.font = _CABECALHO
        celula.fill = _FUNDO
        celula.alignment = Alignment(vertical="center", wrap_text=True)
    for indice, largura in enumerate(larguras, start=1):
        aba.column_dimensions[get_column_letter(indice)].width = largura
    aba.freeze_panes = "A2"
    if aba.max_row > 1:
        aba.auto_filter.ref = aba.dimensions


def exportar_rfqs(base: BaseDados, destino: Path, rfqs: list[RFQ] | None = None, hoje: date | None = None) -> Path:
    rfqs = sorted(rfqs if rfqs is not None else base.rfqs.todos(), key=lambda r: r.numero)
    livro = Workbook()
    aba_rfqs = livro.active
    aba_rfqs.title = "RFQs"
    aba_rfqs.append([
        "RFQ Nº", "Situação", "Data de envio", "Prazo", "Data da resposta", "Projeto", "Cliente", "Planta",
        "Fornecedor", "Idioma", "Pacote", "Nº de itens", "Valor total", "Moeda", "Lead time",
        "Validade da proposta", "Observações",
    ])
    aba_itens = livro.create_sheet("Itens")
    aba_itens.append([
        "RFQ Nº", "Data de envio", "Cliente", "Projeto", "Fornecedor", "OP - Ref", "CDC - Ref",
        "Descrição", "Unidade", "Quantidade", "Volume anual", "Moeda", "Preço unitário", "Ferramental",
        "Lead time (dias)", "Frete", "Embalagem", "Situação",
    ])

    for rfq in rfqs:
        pacote = base.pacotes.obter(rfq.pacote_id)
        projeto = base.projeto_do_pacote(pacote)
        fornecedor = base.fornecedores.obter(rfq.fornecedor_id)
        situacao = rfq.situacao(hoje)
        nome_projeto = projeto.nome if projeto else ""
        cliente = projeto.cliente if projeto else ""
        nome_fornecedor = fornecedor.nome if fornecedor else ""
        aba_rfqs.append([
            rfq.numero, situacao, rfq.data_envio, rfq.prazo, rfq.data_resposta, nome_projeto, cliente,
            projeto.planta if projeto else "", nome_fornecedor, rfq.idioma.value,
            pacote.titulo if pacote else "", len(pacote.itens) if pacote else 0, rfq.valor_total, rfq.moeda,
            rfq.lead_time, rfq.validade, rfq.observacoes,
        ])
        for item in pacote.itens if pacote else []:
            cotacao = rfq.cotacao_do_item(item.id)
            valores = ([cotacao.preco, cotacao.ferramental, cotacao.lead_time_dias, cotacao.frete, cotacao.embalagem]
                       if cotacao else [None] * 5)
            aba_itens.append([
                rfq.numero, rfq.data_envio, cliente, nome_projeto, nome_fornecedor,
                item.ref_op, item.ref_cdc, item.descricao, item.unidade, item.quantidade, item.volume_anual,
                rfq.moeda if cotacao else "", *valores, situacao,
            ])

    for aba, colunas in ((aba_rfqs, "CDEP"), (aba_itens, "B")):
        for coluna in colunas:
            for celula in aba[coluna][1:]:
                celula.number_format = "DD/MM/YYYY"
    for celula in aba_rfqs["M"][1:]:
        celula.number_format = "#,##0.00"
    for coluna in "JKO":
        for celula in aba_itens[coluna][1:]:
            celula.number_format = "#,##0.##"
    for coluna in "MNPQ":
        for celula in aba_itens[coluna][1:]:
            celula.number_format = _FORMATO_VALOR

    _formatar_aba(aba_rfqs, [14, 15, 13, 13, 13, 28, 16, 18, 30, 8, 34, 10, 14, 8, 14, 13, 40])
    _formatar_aba(aba_itens, [14, 13, 16, 28, 30, 36, 44, 30, 10, 13, 14, 8, 14, 14, 10, 12, 12, 15])
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    livro.save(destino)
    return destino


def exportar_comparativo(base: BaseDados, pacote_id: str, destino: Path) -> Path:
    """Comparativo de propostas do pacote em Excel: resumo por fornecedor e itens lado a lado."""
    comparativo = comparar(base, pacote_id)
    moeda = comparativo.moeda
    livro = Workbook()
    resumo = livro.active
    resumo.title = "Resumo"
    resumo.append([
        "RFQ Nº", "Fornecedor", "Situação", "Moeda", f"Câmbio para {moeda}", "Validade", "Itens cotados",
        f"Custo das peças ({moeda})", f"Ferramental ({moeda})", "Maior lead time (dias)",
    ])
    total_itens = len(comparativo.itens)
    for proposta in comparativo.propostas:
        rfq = proposta.rfq
        resumo.append([
            rfq.numero, proposta.fornecedor, rfq.situacao(), proposta.moeda, proposta.taxa, rfq.validade,
            f"{proposta.itens_cotados}/{total_itens}", proposta.total_pecas, proposta.ferramental, proposta.lead_time,
        ])
        linha = resumo.max_row
        for coluna, melhores in (("H", comparativo.melhor_total), ("I", comparativo.melhor_ferramental),
                                 ("J", comparativo.melhor_lead_time)):
            if rfq.id in melhores:
                resumo[f"{coluna}{linha}"].fill = _MELHOR
    for coluna in "EHI":
        for celula in resumo[coluna][1:]:
            celula.number_format = _FORMATO_VALOR
    for celula in resumo["F"][1:]:
        celula.number_format = "DD/MM/YYYY"
    _formatar_aba(resumo, [14, 30, 15, 8, 14, 12, 12, 20, 18, 16])
    resumo.append([])
    resumo.append([f"Pacote: {comparativo.pacote.titulo}"])
    resumo.append([f"Custo das peças = (preço + frete + embalagem) × volume anual (ou quantidade), em {moeda}. "
                   "Verde: melhor valor entre as propostas completas."])

    itens = livro.create_sheet("Itens")
    cabecalho = ["#", "OP - Ref", "CDC - Ref", "Descrição", "Quantidade", "Volume anual"]
    for proposta in comparativo.propostas:
        nome = f"{proposta.fornecedor} ({proposta.rfq.numero})"
        cabecalho += [f"{nome} — custo unit. ({moeda})", f"{nome} — ferramental ({moeda})",
                      f"{nome} — lead time (dias)"]
    itens.append(cabecalho)
    for posicao, item in enumerate(comparativo.itens, start=1):
        valores: list = [posicao, item.ref_op, item.ref_cdc, item.descricao, item.quantidade, item.volume_anual]
        destaques = []
        for proposta in comparativo.propostas:
            cotacao = proposta.item(item.id)
            for campo in ("custo", "ferramental", "lead_time"):
                if cotacao is None:
                    valor = None
                elif campo == "custo":
                    valor = proposta.converter(cotacao.custo_unitario())
                elif campo == "ferramental":
                    valor = proposta.converter(cotacao.ferramental)
                else:
                    valor = cotacao.lead_time_dias
                valores.append(valor)
                destaques.append(comparativo.melhor(item.id, campo, proposta.rfq.id))
        itens.append(valores)
        for deslocamento, melhor in enumerate(destaques):
            celula = itens.cell(row=itens.max_row, column=7 + deslocamento)
            celula.number_format = "#,##0" if deslocamento % 3 == 2 else _FORMATO_VALOR
            if melhor:
                celula.fill = _MELHOR
    _formatar_aba(itens, [5, 24, 30, 26, 12, 14] + [22] * (3 * len(comparativo.propostas)))
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    livro.save(destino)
    return destino
